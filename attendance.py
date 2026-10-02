import streamlit as st
import pandas as pd
from datetime import datetime, timedelta
import os
import base64
import re
import io
import zipfile
import json
import hashlib
from supabase import create_client

st.set_page_config(page_title="ระบบบันทึกเวลาทำงาน", layout="wide")
st.title("📋 ระบบบันทึกเวลาทำงาน")

# ==========================================
# ฐานข้อมูลถาวร Supabase
# ==========================================
def get_secret(name):
    try:
        value = st.secrets.get(name)
        return value
    except Exception:
        return None

SUPABASE_URL = get_secret("SUPABASE_URL")
SUPABASE_KEY = get_secret("SUPABASE_KEY")

if not SUPABASE_URL or not SUPABASE_KEY:
    st.error("ยังไม่ได้เชื่อมต่อฐานข้อมูลถาวรค่ะ")
    st.info("ไปที่ Streamlit Cloud → Settings → Secrets แล้วเพิ่ม SUPABASE_URL และ SUPABASE_KEY จาก Project CLC Crane")
    st.code('SUPABASE_URL = "https://xxxxxxxx.supabase.co"\nSUPABASE_KEY = "your-key"')
    st.stop()

try:
    supabase = create_client(SUPABASE_URL, SUPABASE_KEY)
except Exception as e:
    st.error(f"เชื่อมต่อ Supabase ไม่สำเร็จ: {e}")
    st.stop()

DB_EMP_KEY = "employees"
DB_REC_KEY = "records"
DB_REC_TABLE = "attendance_records"
DB_USER_KEY = "user"

def _json_safe(value):
    return json.loads(json.dumps(value, ensure_ascii=False, default=str))

def db_get(key, default=None):
    try:
        res = supabase.table("app_store").select("value").eq("key", key).limit(1).execute()
        if res.data:
            return res.data[0].get("value", default)
    except Exception as e:
        st.error(f"อ่านข้อมูลฐานข้อมูลไม่สำเร็จ: {e}")
        st.stop()
    return default

def db_set(key, value):
    try:
        supabase.table("app_store").upsert({"key": key, "value": _json_safe(value)}, on_conflict="key").execute()
        return True
    except Exception as e:
        st.error(f"บันทึกข้อมูลฐานข้อมูลไม่สำเร็จ: {e}")
        return False

def load_employees():
    data = db_get(DB_EMP_KEY, None)
    if data is not None:
        return pd.DataFrame(data, columns=["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"])
    # ย้ายข้อมูลเก่าจาก CSV เข้า Supabase ครั้งแรก ถ้ามี
    cols = ["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"]
    if os.path.exists(EMP_FILE) and os.path.getsize(EMP_FILE) > 0:
        try:
            df = pd.read_csv(EMP_FILE, encoding="utf-8")
        except UnicodeDecodeError:
            df = pd.read_csv(EMP_FILE, encoding="cp874")
        for col in cols:
            if col not in df.columns:
                df[col] = ""
        df = df[cols].copy()
    else:
        df = pd.DataFrame(columns=cols)
    db_set(DB_EMP_KEY, df.to_dict(orient="records"))
    return df

def save_employees(df):
    cols = ["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"]
    x = df.copy()
    for col in cols:
        if col not in x.columns:
            x[col] = ""
    x = x[cols].copy()
    # ป้องกันค่า NaN จาก CSV ทำให้ JSON/Supabase บันทึกไม่สำเร็จ
    for col in cols:
        x[col] = x[col].fillna("").astype(str)
    payload = x.to_dict(orient="records")
    ok = db_set(DB_EMP_KEY, payload)
    if ok:
        try:
            x.to_csv(EMP_FILE, index=False, encoding="utf-8")
        except Exception:
            pass
    return ok

def sync_employee_master_fields(df):
    """ให้ข้อมูลชื่อ/ชื่อเล่น/ตำแหน่งในรายการลงเวลาอ้างอิงจากรายชื่อพนักงานปัจจุบันตามรหัสพนักงาน
    เพื่อป้องกันข้อมูลเก่าที่บันทึกไว้ก่อนแก้ไขรายชื่อ ทำให้ชื่อเล่น/ตำแหน่งแสดงผิด
    """
    if df is None or df.empty:
        return df

    x = df.copy()
    emp = load_employees()
    if emp is None or emp.empty:
        return x

    emp_cols = ["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"]
    for col in emp_cols:
        if col not in emp.columns:
            emp[col] = ""

    # ใช้รหัสพนักงานเป็นตัวจับคู่ ไม่ใช้ชื่อ/ชื่อเล่น/ตำแหน่ง
    emp_map = emp[emp_cols].copy()
    emp_map["รหัสพนักงาน"] = emp_map["รหัสพนักงาน"].astype(str).str.strip()
    emp_map = emp_map.drop_duplicates(subset=["รหัสพนักงาน"], keep="last")
    emp_map = emp_map.set_index("รหัสพนักงาน")

    eid = x["รหัสพนักงาน"].astype(str).str.strip()
    for col in ["ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"]:
        mapped = eid.map(emp_map[col])
        # ถ้ามีข้อมูลพนักงานปัจจุบัน ให้ใช้ข้อมูลนั้น
        x[col] = mapped.where(mapped.notna(), x[col])

    return x

def _record_key(row):
    eid = str(row.get("รหัสพนักงาน", "")).strip()
    d = str(row.get("วันที่", "")).strip()
    return f"{eid}|{d}"

def _record_payload(df):
    x = df.copy()
    for col in REC_COLUMNS:
        if col not in x.columns:
            x[col] = pd.NA
    x = x[REC_COLUMNS].copy()
    payload = []
    for _, row in x.iterrows():
        data = {}
        for col in REC_COLUMNS:
            val = row[col]
            if pd.isna(val):
                val = None
            elif isinstance(val, (pd.Timestamp, datetime)):
                val = val.strftime("%Y-%m-%d")
            elif hasattr(val, "item"):
                try:
                    val = val.item()
                except Exception:
                    pass
            data[col] = val
        payload.append({"record_key": _record_key(data), "value": data})
    return payload

def save_records(df):
    """บันทึกแต่ละรายการลงคนละแถวใน Supabase แทนการยัด records ทั้งหมดไว้ใน JSON แถวเดียว"""
    try:
        payload = _record_payload(df)
        wanted = {r["record_key"] for r in payload}
        existing = supabase.table(DB_REC_TABLE).select("record_key").limit(10000).execute().data or []
        old_keys = {r.get("record_key") for r in existing if r.get("record_key")}
        for key in old_keys - wanted:
            supabase.table(DB_REC_TABLE).delete().eq("record_key", key).execute()
        if payload:
            supabase.table(DB_REC_TABLE).upsert(payload, on_conflict="record_key").execute()
        try:
            df.to_csv(REC_FILE, index=False, encoding="utf-8")
        except Exception:
            pass
        return True
    except Exception as e:
        st.error(f"บันทึกข้อมูลลงฐานข้อมูลไม่สำเร็จ: {e}")
        return False

# ==========================================
# ระบบ User / Password (1 user)
# ==========================================
USER_FILE = "user.json"
DEFAULT_USER = {
    "username": "admin",
    "password_hash": hashlib.sha256("admin123".encode("utf-8")).hexdigest()
}

def load_user():
    data = db_get(DB_USER_KEY, None)
    if data and data.get("username") and data.get("password_hash"):
        return data
    # ย้าย User เดิมจาก user.json เข้า Supabase ครั้งแรก ถ้ามี
    if os.path.exists(USER_FILE):
        try:
            with open(USER_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            if data.get("username") and data.get("password_hash"):
                db_set(DB_USER_KEY, data)
                return data
        except Exception:
            pass
    db_set(DB_USER_KEY, DEFAULT_USER)
    return DEFAULT_USER.copy()

def save_user(user):
    db_set(DB_USER_KEY, user)
    try:
        with open(USER_FILE, "w", encoding="utf-8") as f:
            json.dump(user, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

def hash_password(password):
    return hashlib.sha256(password.encode("utf-8")).hexdigest()

user = load_user()

if "logged_in" not in st.session_state:
    st.session_state.logged_in = False
    st.session_state.username = None

if not st.session_state.logged_in:
    st.subheader("🔐 เข้าสู่ระบบ")
    login_user = st.text_input("User", value=user.get("username", "admin"))
    login_pass = st.text_input("Password", type="password")
    if st.button("🔓 เข้าสู่ระบบ", type="primary"):
        if login_user == user.get("username") and hash_password(login_pass) == user.get("password_hash"):
            st.session_state.logged_in = True
            st.session_state.username = login_user
            st.success("เข้าสู่ระบบสำเร็จ 🎉")
            st.balloons()
            st.rerun()
        else:
            st.error("User หรือ Password ไม่ถูกต้อง")
    st.stop()

current_user = st.session_state.username


EMP_FILE = "employees.csv"
REC_FILE = "records.csv"

REC_COLUMNS = [
    "รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง",
    "วันที่", "สถานะ", "เวลาเข้า", "เวลาออก",
    "จำนวนชั่วโมง", "จำนวนวัน WOP", "หมายเหตุ", "เอกสารลาป่วย", "รูปเช็คอิน", "รูปเช็คเอาท์",
    "สายนาที", "สายชม", "OT 1.5(ชม.)", "OT 1(ชม.)"
]

def load_records():
    """อ่านข้อมูลลงเวลาจาก attendance_records ซึ่งแยกเป็นรายรายการ เพื่อไม่ให้ Supabase timeout เมื่อมีรูป"""
    try:
        res = supabase.table(DB_REC_TABLE).select("record_key,value").limit(10000).execute()
        rows = res.data or []
        if rows:
            df = pd.DataFrame([r.get("value", {}) for r in rows])
        else:
            legacy = db_get(DB_REC_KEY, None)
            if legacy:
                df = pd.DataFrame(legacy)
                if not df.empty:
                    save_records(df)
            elif os.path.exists(REC_FILE) and os.path.getsize(REC_FILE) > 0:
                try:
                    df = pd.read_csv(REC_FILE, encoding="utf-8")
                except UnicodeDecodeError:
                    df = pd.read_csv(REC_FILE, encoding="cp874")
                if not df.empty:
                    save_records(df)
            else:
                return pd.DataFrame(columns=REC_COLUMNS)
    except Exception as e:
        st.error(f"อ่านข้อมูลการลงเวลาไม่สำเร็จ: {e}")
        return pd.DataFrame(columns=REC_COLUMNS)

    for col in REC_COLUMNS:
        if col not in df.columns:
            df[col] = pd.NA
    df = df[REC_COLUMNS].copy()
    text_cols = [
        "รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง", "วันที่", "สถานะ",
        "เวลาเข้า", "เวลาออก", "หมายเหตุ", "เอกสารลาป่วย", "รูปเช็คอิน", "รูปเช็คเอาท์"
    ]
    for col in text_cols:
        df[col] = df[col].astype("string")
    numeric_cols = ["จำนวนชั่วโมง", "จำนวนวัน WOP", "สายนาที", "สายชม", "OT 1.5(ชม.)", "OT 1(ชม.)"]
    for col in numeric_cols:
        df[col] = pd.to_numeric(df[col].replace({"-": pd.NA, "": pd.NA}), errors="coerce").astype("float64")
    leave_mask = df["สถานะ"].astype(str).eq("ลากิจไม่รับเงิน")
    for col in ["สายนาที", "สายชม", "OT 1.5(ชม.)", "OT 1(ชม.)"]:
        df.loc[leave_mask, col] = 0.0

    # สำคัญ: ชื่อ/ชื่อเล่น/ตำแหน่งต้องอ้างอิงจากรายชื่อพนักงานปัจจุบัน
    # โดยจับคู่ด้วย "รหัสพนักงาน" เท่านั้น เพื่อแก้ข้อมูลเก่าที่บันทึกชื่อเล่น/ตำแหน่งคลาดเคลื่อน
    df = sync_employee_master_fields(df)

    return df.sort_values("วันที่", key=lambda s: pd.to_datetime(s, errors="coerce"), na_position="last").reset_index(drop=True)


if not os.path.exists(EMP_FILE) or os.path.getsize(EMP_FILE) == 0:
    pd.DataFrame(columns=["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"]).to_csv(EMP_FILE, index=False, encoding="utf-8")
if not os.path.exists(REC_FILE) or os.path.getsize(REC_FILE) == 0:
    pd.DataFrame(columns=[
        "รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง",
        "วันที่", "สถานะ", "เวลาเข้า", "เวลาออก",
        "จำนวนชั่วโมง", "จำนวนวัน WOP", "หมายเหตุ", "เอกสารลาป่วย", "รูปเช็คอิน", "รูปเช็คเอาท์",
        "สายนาที", "สายชม", "OT 1.5(ชม.)", "OT 1(ชม.)"
    ]).to_csv(REC_FILE, index=False, encoding="utf-8")

menu_options = ["บันทึกการเข้างาน/ลา", "หน้าสรุปภาพรวม", "รายงานรายชื่อ+วันทำงาน", "หน้าสรุปส่ง HR", "🔑 เปลี่ยน Password"]
menu_options = ["จัดการรายชื่อพนักงาน"] + menu_options + ["📥 สำรองข้อมูล"]

st.sidebar.success(f"เข้าสู่ระบบ: {current_user} (Admin)")
st.sidebar.caption("☁️ ฐานข้อมูลถาวร: Supabase")
if st.sidebar.button("🚪 ออกจากระบบ"):
    st.session_state.logged_in = False
    st.session_state.username = None
    st.session_state.role = None
    st.rerun()

menu = st.sidebar.selectbox("เมนูหลัก", menu_options)


def flatten_agg_col(c):
    """Flatten pandas groupby/agg column labels without creating duplicate names."""
    if isinstance(c, tuple):
        if len(c) >= 2:
            # Groupby index columns are typically (name, "").
            if c[1] in ("", "sum"):
                return c[0]
            return c[1]
        return c[0]
    return c



def get_payroll_period(base_date=None):
    """Return payroll cycle: 26th through 25th of the following month.
    Example: 26/09/2026 - 25/10/2026.
    """
    if base_date is None:
        base_date = datetime.now().date()
    else:
        base_date = base_date if hasattr(base_date, "year") else datetime.now().date()

    if base_date.day >= 26:
        start_date = base_date.replace(day=26)
        if base_date.month == 12:
            end_date = base_date.replace(year=base_date.year + 1, month=1, day=25)
        else:
            end_date = base_date.replace(month=base_date.month + 1, day=25)
    else:
        if base_date.month == 1:
            start_date = base_date.replace(year=base_date.year - 1, month=12, day=26)
        else:
            start_date = base_date.replace(month=base_date.month - 1, day=26)
        end_date = base_date.replace(day=25)
    return start_date, end_date

def extract_leave_hours_from_note(note):
    """Recover hours from an old note such as 'ลากิจไม่รับเงิน 1 ชม.' when needed."""
    if note is None or pd.isna(note):
        return 0.0
    m = re.search(r"(\d+(?:\.\d+)?)\s*ชม\.?", str(note))
    return float(m.group(1)) if m else 0.0


def add_unpaid_leave_hours(df):
    """สร้างชั่วโมงลากิจไม่รับเงินเป็นคอลัมน์ตัวเลข โดยไม่แก้ค่าใน Series เดิมแบบเสี่ยง dtype error."""
    out = df.copy()
    hours = pd.to_numeric(out["จำนวนชั่วโมง"], errors="coerce").astype("float64")
    leave_mask = out["สถานะ"].astype(str).eq("ลากิจไม่รับเงิน")
    recover_mask = leave_mask & (hours.isna() | (hours <= 0))

    # สร้างค่าที่จะใช้สำหรับทุกแถวก่อน แล้วเลือกด้วย where/mask
    # เพื่อหลีกเลี่ยง LossySetitemError จาก pandas รุ่นใหม่
    recovered_all = pd.Series(0.0, index=out.index, dtype="float64")
    if "หมายเหตุ" in out.columns:
        recovered_all = out["หมายเหตุ"].apply(extract_leave_hours_from_note)
        recovered_all = pd.to_numeric(recovered_all, errors="coerce").fillna(0.0).astype("float64")

    # ถ้ามีจำนวนชั่วโมงจริง > 0 ให้ใช้ค่านั้น; ถ้าไม่มี ให้ใช้ค่าจากหมายเหตุ
    final_hours = hours.where(~recover_mask, recovered_all)
    out["_ลากิจไม่รับเงินชม"] = final_hours.fillna(0.0).where(leave_mask, 0.0).astype("float64")
    return out




def exclude_post_resignation_records(df):
    """ไม่แสดง/ไม่นับข้อมูลหลังวันลาออกของพนักงานคนนั้นในรายงานและสรุป"""
    if df is None or df.empty:
        return df
    x = df.copy()
    x["วันที่"] = pd.to_datetime(x["วันที่"], errors="coerce")
    resign = x[x["สถานะ"].astype(str).eq("ลาออก")][["รหัสพนักงาน", "วันที่"]].dropna().copy()
    if resign.empty:
        return x
    resign = resign.groupby("รหัสพนักงาน")["วันที่"].min().to_dict()
    keep = []
    for _, row in x.iterrows():
        rid = str(row.get("รหัสพนักงาน", ""))
        d = row.get("วันที่")
        rdate = resign.get(rid)
        keep.append(not (rdate is not None and pd.notna(d) and d > rdate))
    return x.loc[keep].copy()


def count_checked_days(df, start_d=None, end_d=None):
    """นับวันทำงานจริงจากวันที่มีทั้งเวลาเข้าและเวลาออก ไม่ขึ้นกับวันในสัปดาห์."""
    if df is None or df.empty:
        return 0
    x = df.copy()
    x["วันที่"] = pd.to_datetime(x["วันที่"], errors="coerce")
    if start_d is not None:
        x = x[x["วันที่"] >= pd.to_datetime(start_d)]
    if end_d is not None:
        x = x[x["วันที่"] <= pd.to_datetime(end_d)]
    if x.empty:
        return 0
    tin = x["เวลาเข้า"].astype(str).str.strip()
    tout = x["เวลาออก"].astype(str).str.strip()
    valid = (tin.notna() & tout.notna() & tin.ne("") & tout.ne("") &
             ~tin.isin(["-", "nan", "NaT"]) & ~tout.isin(["-", "nan", "NaT"]))
    return int(x.loc[valid, "วันที่"].dropna().dt.normalize().nunique())


def get_employee_resignation_date(df, employee_id):
    """คืนวันที่ลาออกของพนักงาน หากมีการบันทึกสถานะลาออกแล้ว"""
    if df is None or df.empty:
        return None
    mask = (df["รหัสพนักงาน"].astype(str) == str(employee_id)) & (df["สถานะ"].astype(str) == "ลาออก")
    dates = pd.to_datetime(df.loc[mask, "วันที่"], errors="coerce").dropna()
    return dates.min().date() if not dates.empty else None


def checked_days_by_employee(df):
    """คืนจำนวนวันที่มีเวลาเข้าและเวลาออกจริง แยกตามพนักงาน."""
    keys = ["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"]
    if df is None or df.empty:
        return pd.DataFrame(columns=keys + ["วันทำงานทั้งหมด"])
    x = df.copy()
    x["วันที่"] = pd.to_datetime(x["วันที่"], errors="coerce")
    tin = x["เวลาเข้า"].astype(str).str.strip()
    tout = x["เวลาออก"].astype(str).str.strip()
    valid = (tin.notna() & tout.notna() & tin.ne("") & tout.ne("") &
             ~tin.isin(["-", "nan", "NaT"]) & ~tout.isin(["-", "nan", "NaT"]))
    x = x.loc[valid & x["วันที่"].notna()].copy()
    if x.empty:
        return x[keys].drop_duplicates().assign(**{"วันทำงานทั้งหมด": 0})
    x["วันที่"] = x["วันที่"].dt.normalize()
    out = x.groupby(keys, dropna=False)["วันที่"].nunique().reset_index(name="วันทำงานทั้งหมด")
    return out


def order_summary_columns(df):
    """จัดลำดับคอลัมน์สรุป โดยให้วันทำงานทั้งหมดอยู่ก่อนมาปกติ(วัน)."""
    wanted = [
        "รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง",
        "วันทำงานทั้งหมด", "มาปกติ(วัน)", "ลาป่วย(วัน)", "WOP(วัน)", "ลาออก(วัน)",
        "ลากิจไม่รับเงิน(ชม.)", "ขาดงาน(วัน)",
        "OT 1.5(ชม.)", "OT 1(ชม.)", "สายชม", "สายนาที",
        "มาสายรวม(ชม.)", "มาสายรวม(นาที)"
    ]
    cols = list(df.columns)
    first = [c for c in wanted if c in cols]
    rest = [c for c in cols if c not in first]
    return df[first + rest]


def order_hr_columns(df):
    """คงวันที่ไว้หลังข้อมูลพนักงาน แล้ววางสรุปตามลำดับที่กำหนด."""
    emp = ["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"]
    summary = [
        "วันทำงานทั้งหมด", "มาปกติ(วัน)", "ลาป่วย(วัน)", "WOP(วัน)", "ลาออก(วัน)",
        "ลากิจไม่รับเงิน(ชม.)", "ขาดงาน(วัน)",
        "OT 1.5(ชม.)", "OT 1(ชม.)", "มาสายรวม(ชม.)", "มาสายรวม(นาที)"
    ]
    cols = list(df.columns)
    date_cols = [c for c in cols if c not in emp + summary]
    return df[[c for c in emp if c in cols] + date_cols + [c for c in summary if c in cols]]

def img_cell(val):
    s = str(val)
    if s.startswith("data:image"):
        return f'<img src="{s}" width="70" />'
    return "-"

def parse_time_str(s):
    try:
        return datetime.strptime(str(s), "%H:%M")
    except:
        return None

# ==========================================
# เปลี่ยน Password
# ==========================================
if menu == "🔑 เปลี่ยน Password":
    st.header("🔑 เปลี่ยน Password")
    st.write(f"User: **{current_user}**")
    with st.form("change_password_form"):
        new_user = st.text_input("User ใหม่", value=user.get("username", "admin"))
        old_pw = st.text_input("Password เดิม", type="password")
        new_pw = st.text_input("Password ใหม่", type="password")
        confirm_pw = st.text_input("ยืนยัน Password ใหม่", type="password")
        submit_pw = st.form_submit_button("💾 บันทึก User / Password")
    if submit_pw:
        if hash_password(old_pw) != user.get("password_hash"):
            st.error("Password เดิมไม่ถูกต้อง")
        elif not new_user.strip():
            st.error("กรุณาระบุ User")
        elif len(new_pw) < 4:
            st.error("Password ใหม่ต้องมีอย่างน้อย 4 ตัวอักษร")
        elif new_pw != confirm_pw:
            st.error("Password ใหม่ไม่ตรงกัน")
        else:
            user["username"] = new_user.strip()
            user["password_hash"] = hash_password(new_pw)
            save_user(user)
            st.session_state.username = user["username"]
            st.success("เปลี่ยน User / Password สำเร็จ 🎉")
            st.balloons()

# ==========================================
# 1. จัดการรายชื่อพนักงาน
# ==========================================
if menu == "จัดการรายชื่อพนักงาน":
    st.header("จัดการรายชื่อพนักงาน")
    df_emp = load_employees()
    
    c1, c2, c3, c4 = st.columns(4)
    with c1: eid = st.text_input("รหัสพนักงาน")
    with c2: enam = st.text_input("ชื่อพนักงาน")
    with c3: enick = st.text_input("ชื่อเล่น")
    with c4: epos = st.text_input("ตำแหน่ง")
    
    if st.button("เพิ่มพนักงาน") and eid and enam:
        exists = False
        if "รหัสพนักงาน" in df_emp.columns:
            exists = eid in df_emp["รหัสพนักงาน"].astype(str).values
        if not exists:
            new = pd.DataFrame([[eid, enam, enick, epos]],
                columns=["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"])
            df_emp = pd.concat([df_emp, new], ignore_index=True)
            save_employees(df_emp)
            st.success("เพิ่มพนักงานสำเร็จ 🎉")
            st.balloons()
            st.rerun()
        else:
            st.warning("มีรหัสนี้อยู่แล้ว ⚠️")
    
    st.subheader("รายชื่อทั้งหมด")
    if df_emp.empty:
        st.info("ยังไม่มีข้อมูล")
    else:
        if "edit_idx" not in st.session_state:
            st.session_state.edit_idx = None
        for idx, row in df_emp.iterrows():
            eid_row = row["รหัสพนักงาน"]
            name_row = row["ชื่อพนักงาน"]
            nick_row = row["ชื่อเล่น"]
            pos_row = row["ตำแหน่ง"]
            if st.session_state.edit_idx == idx:
                with st.form(f"edit_form_{idx}"):
                    ceid = st.text_input("รหัสพนักงาน", value=eid_row, disabled=True)
                    cnam = st.text_input("ชื่อพนักงาน", value=name_row)
                    cnick = st.text_input("ชื่อเล่น", value=nick_row)
                    cpos = st.text_input("ตำแหน่ง", value=pos_row)
                    sv, cl = st.columns(2)
                    with sv:
                        if st.form_submit_button("💾 บันทึก"):
                            df_emp.at[idx, "ชื่อพนักงาน"] = cnam
                            df_emp.at[idx, "ชื่อเล่น"] = cnick
                            df_emp.at[idx, "ตำแหน่ง"] = cpos
                            save_employees(df_emp)
                            st.session_state.edit_idx = None
                            st.success("บันทึกการแก้ไขสำเร็จ 🎉")
                            st.balloons()
                            st.rerun()
                    with cl:
                        if st.form_submit_button("❌ ยกเลิก"):
                            st.session_state.edit_idx = None
                            st.rerun()
            else:
                c1, c2, c3, c4, c5, c6 = st.columns([2, 3, 2, 2, 1.2, 1.2])
                c1.write(eid_row)
                c2.write(name_row)
                c3.write(nick_row)
                c4.write(pos_row)
                if c5.button("✏️ แก้ไข", key=f"edit_btn_{idx}"):
                    st.session_state.edit_idx = idx
                    st.rerun()
                if c6.button("🗑️ ลบ", key=f"del_btn_{idx}"):
                    st.session_state[f"confirm_del_{idx}"] = True
                if f"confirm_del_{idx}" in st.session_state and st.session_state[f"confirm_del_{idx}"]:
                    st.warning(f"ต้องการลบ {name_row} ใช่หรือไม่?")
                    y, n = st.columns(2)
                    if y.button("✅ ยืนยันลบ", key=f"del_ok_{idx}"):
                        df_emp = df_emp.drop(idx).reset_index(drop=True)
                        save_employees(df_emp)
                        del st.session_state[f"confirm_del_{idx}"]
                        st.success("ลบสำเร็จ ✅")
                        st.rerun()
                    if n.button("❌ ยกเลิก", key=f"del_no_{idx}"):
                        del st.session_state[f"confirm_del_{idx}"]

# ==========================================
# 2. บันทึกการเข้างาน/ลา
# ==========================================
elif menu == "บันทึกการเข้างาน/ลา":
    st.header("บันทึกการเข้างาน/ลา")
    df_emp = load_employees()
    if df_emp.empty:
        st.warning("กรุณาเพิ่มรายชื่อพนักงานก่อน ⚠️")
        st.stop()

    sel = st.selectbox("เลือกพนักงาน", [
        f"{r['รหัสพนักงาน']} | {r['ชื่อพนักงาน']} — {r['ตำแหน่ง']}"
        for _, r in df_emp.iterrows()
    ])
    eid = sel.split(" | ")[0]
    enam = sel.split(" | ")[1].split(" — ")[0]
    emp_row = df_emp[df_emp["รหัสพนักงาน"].astype(str) == eid].iloc[0]
    enick = emp_row["ชื่อเล่น"]
    epos = emp_row["ตำแหน่ง"]

    # โหลดข้อมูลก่อนเลือกวันที่ เพื่อป้องกันการคีย์วันที่หลังวันลาออกของพนักงานคนนั้น
    df_rec = load_records()
    resignation_date = get_employee_resignation_date(df_rec, eid)
    default_entry_date = datetime.now().date()
    if resignation_date is not None and default_entry_date > resignation_date:
        default_entry_date = resignation_date

    if resignation_date is not None:
        st.info(f"🚪 พนักงานรายนี้ลาออกวันที่ {resignation_date.strftime('%d/%m/%y')} — วันที่หลังจากนี้จะไม่สามารถคีย์ได้")
        dt = st.date_input("วันที่", default_entry_date, max_value=resignation_date, key=f"work_date_{eid}")
        if dt > resignation_date:
            st.error(f"⛔ ไม่สามารถคีย์วันที่ {dt.strftime('%d/%m/%y')} ได้ เพราะพนักงานลาออกวันที่ {resignation_date.strftime('%d/%m/%y')}")
            st.stop()
    else:
        dt = st.date_input("วันที่", default_entry_date, key=f"work_date_{eid}")

    stt = st.selectbox("สถานะ", [
        "มาปกติ", "ลาป่วย", "ลากิจไม่รับเงิน", "WOP", "ขาดงาน", "ลาออก"
    ])

    today_mask = (
        (df_rec["รหัสพนักงาน"].astype(str) == str(eid)) &
        (pd.to_datetime(df_rec["วันที่"], errors="coerce").dt.date == dt)
    )
    existing = df_rec.loc[today_mask]

    tin = tout = pin = pout = note = sick_doc = ""
    leave_hours_existing = 0.0
    wop_days_existing = 1.0
    late_min = late_hr = ot15 = ot1 = hrs_work = 0.0
    if not existing.empty:
        rec = existing.iloc[0]
        tin = str(rec["เวลาเข้า"]) if pd.notna(rec["เวลาเข้า"]) and rec["เวลาเข้า"] != "-" else ""
        tout = str(rec["เวลาออก"]) if pd.notna(rec["เวลาออก"]) and rec["เวลาออก"] != "-" else ""
        pin = str(rec["รูปเช็คอิน"]) if pd.notna(rec["รูปเช็คอิน"]) and rec["รูปเช็คอิน"] != "-" else ""
        pout = str(rec["รูปเช็คเอาท์"]) if pd.notna(rec["รูปเช็คเอาท์"]) and rec["รูปเช็คเอาท์"] != "-" else ""
        note = str(rec["หมายเหตุ"]) if pd.notna(rec["หมายเหตุ"]) and rec["หมายเหตุ"] != "-" else ""
        sick_doc = str(rec["เอกสารลาป่วย"]) if "เอกสารลาป่วย" in rec.index and pd.notna(rec["เอกสารลาป่วย"]) and rec["เอกสารลาป่วย"] != "-" else ""
        if "จำนวนวัน WOP" in rec.index and pd.notna(rec["จำนวนวัน WOP"]):
            try:
                wop_days_existing = float(rec["จำนวนวัน WOP"])
            except Exception:
                wop_days_existing = 1.0
        leave_hours_existing = (
            float(rec["จำนวนชั่วโมง"])
            if pd.notna(rec["จำนวนชั่วโมง"])
            else extract_leave_hours_from_note(note)
        )
        st.info(f"📋 มีบันทึกของวันนี้แล้ว | เวลาเข้า: {tin or '—'} | เวลาออก: {tout or '—'}")

    def calc_times(tin_str, tout_str, work_date):
        t_in = parse_time_str(tin_str)
        t_out = parse_time_str(tout_str)
        if not t_in or not t_out:
            return 0, 0, 0, 0, 0.0
        ref_8am = t_in.replace(hour=8, minute=0, second=0)
        late_delta = (t_in - ref_8am).total_seconds() / 60 if t_in > ref_8am else 0
        late_hr_val = round(late_delta / 60, 2) if late_delta >= 60 else 0
        late_min_val = late_delta % 60 if late_delta > 0 else 0
        day_of_week = work_date.weekday()
        total_sec = (t_out - t_in).total_seconds() - 3600
        total_hr = round(total_sec / 3600, 2) if total_sec > 0 else 0
        std_hr = 8
        ot15_val = 0.0
        ot1_val = 0.0
        # กติกา OT:
        # - จันทร์-เสาร์: เวลาทำงานปกติ 8 ชม. แล้วเวลาที่เกิน 8 ชม. = OT 1.5
        # - วันอาทิตย์: เวลาที่มาทำงานทั้งหมด = OT 1
        # - ใช้เฉพาะชั่วโมงเต็มของ OT ตามกติกาที่กำหนด (ไม่ปัดเศษเป็นเศษชั่วโมง)
        if day_of_week == 6:
            ot1_val = float(int(total_hr)) if total_hr > 0 else 0.0
        elif total_hr > std_hr:
            ot15_val = float(int(total_hr - std_hr))
        return late_min_val, late_hr_val, ot15_val, ot1_val, total_hr

    # ------------------------------------------------------------
    # มาปกติ / ลากิจไม่รับเงิน(ชม.)
    # ทั้ง 2 สถานะสามารถแก้เวลาเข้า-ออกและแนบรูปเข้า-ออกได้
    # การทำงานล่วงเวลาไม่ใช่สถานะใหม่: ยังเลือก "มาปกติ" และระบบคำนวณ OT อัตโนมัติ
    # ------------------------------------------------------------
    if stt in ["มาปกติ", "ลากิจไม่รับเงิน"]:
        if stt == "ลากิจไม่รับเงิน":
            leave_hours = st.number_input(
                "จำนวนชั่วโมงลากิจไม่รับเงิน",
                min_value=0.0,
                max_value=24.0,
                value=float(leave_hours_existing),
                step=0.5,
                format="%.2f",
                key=f"leave_hours_{eid}_{dt}"
            )
            st.caption("💡 กรณีลากิจไม่รับเงิน 1 ชั่วโมง เช่น 08:00–09:00 ให้กรอก 1.00 ชั่วโมง")

        col1, col2 = st.columns(2)
        with col1:
            st.subheader("🕘 เช็คอินเวลาเข้า")
            default_tin = datetime.strptime(tin, "%H:%M") if tin else datetime.strptime("08:00", "%H:%M")
            tin_new = st.time_input("เวลาเข้า", default_tin, key=f"tin_{eid}_{dt}_{stt}").strftime("%H:%M")
            st.write("แนบรูปเช็คอิน")
            fin = st.file_uploader("อัปโหลดรูปเช็คอิน", type=["jpg", "jpeg", "png"], key=f"pin_upload_{eid}_{dt}_{stt}")
            pin_new = f"data:image/jpeg;base64,{base64.b64encode(fin.read()).decode()}" if fin else pin

            if st.button("✅ บันทึกเวลาเข้า", type="primary", key=f"save_in_{eid}_{dt}_{stt}"):
                if stt == "ลากิจไม่รับเงิน":
                    lmin, lhr, ot15, ot1, wh = 0.0, 0.0, 0.0, 0.0, 0.0
                    qty = float(leave_hours)
                else:
                    lmin, lhr, ot15, ot1, wh = calc_times(tin_new, tout, dt)
                    qty = wh if tout else None
                if existing.empty:
                    new_rec = pd.DataFrame([[
                        eid, enam, enick, epos, dt, stt, tin_new, tout,
                        qty, 0.0, note, sick_doc, pin_new, pout, lmin, lhr, ot15, ot1
                    ]], columns=REC_COLUMNS)
                    df_rec = pd.concat([df_rec, new_rec], ignore_index=True)
                else:
                    idx = existing.index[0]
                    df_rec.at[idx, "สถานะ"] = stt
                    df_rec.at[idx, "เวลาเข้า"] = tin_new
                    df_rec.at[idx, "รูปเช็คอิน"] = pin_new
                    if stt == "ลากิจไม่รับเงิน":
                        df_rec.at[idx, "จำนวนชั่วโมง"] = float(leave_hours)
                        df_rec.at[idx, "สายนาที"] = 0.0
                        df_rec.at[idx, "สายชม"] = 0.0
                        df_rec.at[idx, "OT 1.5(ชม.)"] = 0.0
                        df_rec.at[idx, "OT 1(ชม.)"] = 0.0
                    elif tout:
                        df_rec.at[idx, "จำนวนชั่วโมง"] = wh
                        df_rec.at[idx, "สายนาที"] = lmin
                        df_rec.at[idx, "สายชม"] = lhr
                        df_rec.at[idx, "OT 1.5(ชม.)"] = ot15
                        df_rec.at[idx, "OT 1(ชม.)"] = ot1
                save_records(df_rec)
                st.success("บันทึกเวลาเข้าสำเร็จ 🎉")
                st.balloons()
                st.rerun()

        with col2:
            st.subheader("🕕 เช็คเอาท์เวลาออก")
            default_tout = datetime.strptime(tout, "%H:%M") if tout else datetime.strptime("17:00", "%H:%M")
            tout_new = st.time_input("เวลาออก", default_tout, key=f"tout_{eid}_{dt}_{stt}").strftime("%H:%M")
            st.write("แนบรูปเช็คเอาท์")
            fout = st.file_uploader("อัปโหลดรูปเช็คเอาท์", type=["jpg", "jpeg", "png"], key=f"pout_upload_{eid}_{dt}_{stt}")
            pout_new = f"data:image/jpeg;base64,{base64.b64encode(fout.read()).decode()}" if fout else pout

            if st.button("✅ บันทึกเวลาออก", type="primary", key=f"save_out_{eid}_{dt}_{stt}"):
                if stt == "ลากิจไม่รับเงิน":
                    lmin, lhr, ot15, ot1, wh = 0.0, 0.0, 0.0, 0.0, 0.0
                    qty = float(leave_hours)
                else:
                    lmin, lhr, ot15, ot1, wh = calc_times(tin or tin_new, tout_new, dt)
                    qty = wh if tin else None
                if existing.empty:
                    new_rec = pd.DataFrame([[
                        eid, enam, enick, epos, dt, stt, tin or tin_new, tout_new,
                        qty, 0.0, note, sick_doc, pin, pout_new, lmin, lhr, ot15, ot1
                    ]], columns=REC_COLUMNS)
                    df_rec = pd.concat([df_rec, new_rec], ignore_index=True)
                else:
                    idx = existing.index[0]
                    df_rec.at[idx, "สถานะ"] = stt
                    df_rec.at[idx, "เวลาออก"] = tout_new
                    df_rec.at[idx, "รูปเช็คเอาท์"] = pout_new
                    if stt == "ลากิจไม่รับเงิน":
                        df_rec.at[idx, "จำนวนชั่วโมง"] = float(leave_hours)
                        df_rec.at[idx, "สายนาที"] = 0.0
                        df_rec.at[idx, "สายชม"] = 0.0
                        df_rec.at[idx, "OT 1.5(ชม.)"] = 0.0
                        df_rec.at[idx, "OT 1(ชม.)"] = 0.0
                    elif tin:
                        df_rec.at[idx, "จำนวนชั่วโมง"] = wh
                        df_rec.at[idx, "สายนาที"] = lmin
                        df_rec.at[idx, "สายชม"] = lhr
                        df_rec.at[idx, "OT 1.5(ชม.)"] = ot15
                        df_rec.at[idx, "OT 1(ชม.)"] = ot1
                save_records(df_rec)
                st.success("บันทึกเวลาออกสำเร็จ 🎉")
                st.balloons()
                st.rerun()

    elif stt == "ลาป่วย":
        st.subheader("🤒 ข้อมูลลาป่วย")
        st.caption("สามารถแนบใบรับรอง/เอกสารลาป่วย หรือกรอกหมายเหตุ/เหตุผลได้")
        sick_file = st.file_uploader(
            "📎 แนบเอกสารลาป่วย (PDF / JPG / JPEG / PNG)",
            type=["pdf", "jpg", "jpeg", "png"],
            key=f"sick_doc_{eid}_{dt}"
        )
        sick_doc_new = sick_doc
        if sick_file:
            raw = sick_file.read()
            mime = sick_file.type or "application/octet-stream"
            sick_doc_new = f"data:{mime};base64,{base64.b64encode(raw).decode()}"
            st.success("แนบเอกสารแล้ว ✅")
        note_new = st.text_area("หมายเหตุ / เหตุผลการลาป่วย", value=note, key=f"sick_note_{eid}_{dt}")
        if sick_doc:
            st.info("📎 มีเอกสารลาป่วยเดิมอยู่แล้ว หากไม่เลือกไฟล์ใหม่ ระบบจะเก็บเอกสารเดิมไว้")
        if st.button("✅ บันทึกข้อมูลลาป่วย", type="primary"):
            if existing.empty:
                new_rec = pd.DataFrame([[
                    eid, enam, enick, epos, dt, stt, "-", "-", None, 0.0,
                    note_new, sick_doc_new, "-", "-", 0, 0, 0, 0
                ]], columns=REC_COLUMNS)
                df_rec = pd.concat([df_rec, new_rec], ignore_index=True)
            else:
                idx = existing.index[0]
                df_rec.at[idx, "สถานะ"] = stt
                df_rec.at[idx, "หมายเหตุ"] = note_new
                df_rec.at[idx, "เอกสารลาป่วย"] = sick_doc_new or "-"
                df_rec.at[idx, "จำนวนชั่วโมง"] = 0.0
                df_rec.at[idx, "จำนวนวัน WOP"] = 0.0
                df_rec.at[idx, "เวลาเข้า"] = "-"
                df_rec.at[idx, "เวลาออก"] = "-"
                df_rec.at[idx, "OT 1.5(ชม.)"] = 0.0
                df_rec.at[idx, "OT 1(ชม.)"] = 0.0
                df_rec.at[idx, "สายชม"] = 0.0
                df_rec.at[idx, "สายนาที"] = 0.0
            save_records(df_rec)
            st.success("บันทึกลาป่วยสำเร็จ 🎉")
            st.balloons()
            st.rerun()

    elif stt == "WOP":
        st.subheader("📅 WOP")
        wop_days = st.number_input(
            "จำนวนวัน WOP",
            min_value=0.5,
            max_value=31.0,
            value=float(wop_days_existing if wop_days_existing > 0 else 1.0),
            step=0.5,
            format="%.1f",
            key=f"wop_days_{eid}_{dt}"
        )
        st.caption("💡 ระบบจะนำจำนวนวัน WOP ไปคำนวณในสรุป HR โดยอัตโนมัติ")
        note_new = st.text_area("หมายเหตุ / เหตุผล", value=note, key=f"wop_note_{eid}_{dt}")
        if st.button("✅ บันทึก WOP", type="primary"):
            if existing.empty:
                new_rec = pd.DataFrame([[
                    eid, enam, enick, epos, dt, stt, "-", "-", None, float(wop_days),
                    note_new, "-", "-", "-", 0, 0, 0, 0
                ]], columns=REC_COLUMNS)
                df_rec = pd.concat([df_rec, new_rec], ignore_index=True)
            else:
                idx = existing.index[0]
                df_rec.at[idx, "สถานะ"] = stt
                df_rec.at[idx, "หมายเหตุ"] = note_new
                df_rec.at[idx, "จำนวนวัน WOP"] = float(wop_days)
                df_rec.at[idx, "จำนวนชั่วโมง"] = 0.0
                df_rec.at[idx, "เวลาเข้า"] = "-"
                df_rec.at[idx, "เวลาออก"] = "-"
                df_rec.at[idx, "OT 1.5(ชม.)"] = 0.0
                df_rec.at[idx, "OT 1(ชม.)"] = 0.0
                df_rec.at[idx, "สายชม"] = 0.0
                df_rec.at[idx, "สายนาที"] = 0.0
            save_records(df_rec)
            st.success("บันทึก WOP สำเร็จ 🎉")
            st.balloons()
            st.rerun()

    else:
        if stt == "ลาออก":
            st.subheader("🚪 ข้อมูลลาออก")
            st.caption("ใช้สถานะนี้สำหรับวันที่พนักงานลาออก เพื่อให้แสดงในสรุปเป็น ลาออก(วัน)")
            note_new = st.text_area("หมายเหตุ / เหตุผลการลาออก", value=note, key=f"resign_note_{eid}_{dt}")
            save_label = "✅ บันทึกการลาออก"
        else:
            note_new = st.text_area("หมายเหตุ / เหตุผล", value=note, key=f"other_note_{eid}_{dt}")
            save_label = "✅ บันทึกข้อมูล"
        if st.button(save_label, type="primary"):
            if existing.empty:
                new_rec = pd.DataFrame([[
                    eid, enam, enick, epos, dt, stt, "-", "-", None, 0.0,
                    note_new, "-", "-", "-", 0, 0, 0, 0
                ]], columns=REC_COLUMNS)
                df_rec = pd.concat([df_rec, new_rec], ignore_index=True)
            else:
                idx = existing.index[0]
                df_rec.at[idx, "สถานะ"] = stt
                df_rec.at[idx, "หมายเหตุ"] = note_new
                df_rec.at[idx, "เอกสารลาป่วย"] = "-"
                df_rec.at[idx, "จำนวนชั่วโมง"] = 0.0
                df_rec.at[idx, "จำนวนวัน WOP"] = 0.0
                df_rec.at[idx, "เวลาเข้า"] = "-"
                df_rec.at[idx, "เวลาออก"] = "-"
                df_rec.at[idx, "OT 1.5(ชม.)"] = 0.0
                df_rec.at[idx, "OT 1(ชม.)"] = 0.0
                df_rec.at[idx, "สายชม"] = 0.0
                df_rec.at[idx, "สายนาที"] = 0.0
            save_records(df_rec)
            st.success("บันทึกสำเร็จ 🎉")
            st.balloons()
            st.rerun()

# ==========================================
# 3. หน้าสรุปภาพรวม
# ==========================================
elif menu == "หน้าสรุปภาพรวม":
    st.header("สรุปภาพรวม")
    df_rec = load_records()
    if df_rec.empty:
        st.info("ยังไม่มีข้อมูล")
        st.stop()
    df_rec["วันที่"] = pd.to_datetime(df_rec["วันที่"], errors="coerce")
    df_rec = exclude_post_resignation_records(df_rec)
    df_rec = add_unpaid_leave_hours(df_rec)
    summary = df_rec.groupby(
        ["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"], dropna=False
    ).agg({
        "สถานะ": [
            # "มาปกติ" และข้อมูลเก่า "ทำงานล่วงเวลา" ให้นับเป็นวันทำงานเหมือนเดิม
            ("มาปกติ(วัน)", lambda x: x.isin(["มาปกติ", "ทำงานล่วงเวลา"]).sum()),
            ("ลาป่วย(วัน)", lambda x: (x=="ลาป่วย").sum()),
            ("WOP(วัน)", lambda x: (x=="WOP").sum()),
            ("ลาออก(วัน)", lambda x: (x=="ลาออก").sum()),
            ("ขาดงาน(วัน)", lambda x: (x=="ขาดงาน").sum()),
        ],
        "OT 1.5(ชม.)": "sum",
        "OT 1(ชม.)": "sum",
        "สายชม": "sum",
        "สายนาที": "sum",
        "_ลากิจไม่รับเงินชม": "sum"
    }).reset_index()
    summary.columns = [flatten_agg_col(c) for c in summary.columns]
    if "จำนวนวัน WOP" in df_rec.columns:
        wop_sum = df_rec.groupby(["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"], dropna=False)["จำนวนวัน WOP"].sum().reset_index()
        summary = summary.drop(columns=["WOP(วัน)"], errors="ignore").merge(wop_sum, on=["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"], how="left")
        summary = summary.rename(columns={"จำนวนวัน WOP": "WOP(วัน)"})
    if "_ลากิจไม่รับเงินชม" in summary.columns:
        summary = summary.rename(columns={"_ลากิจไม่รับเงินชม": "ลากิจไม่รับเงิน(ชม.)"})
    summary = summary.loc[:, ~summary.columns.duplicated()]
    overview_start, overview_end = get_payroll_period()
    checked = checked_days_by_employee(df_rec)
    summary = summary.drop(columns=["วันทำงานทั้งหมด"], errors="ignore").merge(checked, on=["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"], how="left")
    summary["วันทำงานทั้งหมด"] = pd.to_numeric(summary["วันทำงานทั้งหมด"], errors="coerce").fillna(0).astype(int)
    summary = order_summary_columns(summary)
    st.dataframe(summary, use_container_width=True)

# ==========================================
# 4. รายงานรายชื่อ+วันทำงาน
# ==========================================
elif menu == "รายงานรายชื่อ+วันทำงาน":
    st.header("รายงานรายชื่อ+วันทำงาน")
    default_start, default_end = get_payroll_period()
    st.info("📅 รอบเงินเดือนอัตโนมัติ: วันที่ 26 ถึงวันที่ 25 ของเดือนถัดไป (สามารถเปลี่ยนวันที่เองได้)")
    c1, c2 = st.columns(2)
    start_d = c1.date_input("วันที่เริ่ม", default_start)
    end_d = c2.date_input("วันที่สิ้นสุด", default_end)
    st.subheader(f"รอบ {start_d.strftime('%d/%m/%y')} - {end_d.strftime('%d/%m/%y')}")
    if start_d > end_d:
        st.error("วันที่เริ่มต้องไม่เกินวันที่สิ้นสุด")
        st.stop()
    df_rec = load_records()
    if df_rec.empty:
        st.info("ยังไม่มีข้อมูล")
        st.stop()
    df_rec["วันที่"] = pd.to_datetime(df_rec["วันที่"], errors="coerce")
    df_rec = exclude_post_resignation_records(df_rec)
    mask = (df_rec["วันที่"] >= pd.to_datetime(start_d)) & (df_rec["วันที่"] <= pd.to_datetime(end_d))
    df = df_rec.loc[mask].copy()
    # เรียงวันที่ตามจริงเสมอ ไม่ว่าข้อมูลจะถูกคีย์เข้ามาในลำดับใด
    df = df.sort_values(["วันที่", "รหัสพนักงาน"], ascending=[True, True], na_position="last")
    if df.empty:
        st.info("ไม่มีข้อมูลในช่วงวันที่ที่เลือก")
        st.stop()
    st.subheader("สรุปรวม")
    df = add_unpaid_leave_hours(df)
    sum_df = df.groupby(
        ["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"], dropna=False
    ).agg({
        "สถานะ": [
            # "มาปกติ" และข้อมูลเก่า "ทำงานล่วงเวลา" ให้นับเป็นวันทำงานเหมือนเดิม
            ("มาปกติ(วัน)", lambda x: x.isin(["มาปกติ", "ทำงานล่วงเวลา"]).sum()),
            ("ลาป่วย(วัน)", lambda x: (x=="ลาป่วย").sum()),
            ("WOP(วัน)", lambda x: (x=="WOP").sum()),
            ("ลาออก(วัน)", lambda x: (x=="ลาออก").sum()),
            ("ขาดงาน(วัน)", lambda x: (x=="ขาดงาน").sum()),
        ],
        "OT 1.5(ชม.)": "sum",
        "OT 1(ชม.)": "sum",
        "สายชม": "sum",
        "สายนาที": "sum",
        "_ลากิจไม่รับเงินชม": "sum"
    }).reset_index()
    sum_df.columns = [flatten_agg_col(c) for c in sum_df.columns]
    if "จำนวนวัน WOP" in df.columns:
        wop_sum = df.groupby(["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"], dropna=False)["จำนวนวัน WOP"].sum().reset_index()
        sum_df = sum_df.drop(columns=["WOP(วัน)"], errors="ignore").merge(wop_sum, on=["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"], how="left")
        sum_df = sum_df.rename(columns={"จำนวนวัน WOP": "WOP(วัน)"})
    if "_ลากิจไม่รับเงินชม" in sum_df.columns:
        sum_df = sum_df.rename(columns={"_ลากิจไม่รับเงินชม": "ลากิจไม่รับเงิน(ชม.)"})
    sum_df = sum_df.loc[:, ~sum_df.columns.duplicated()]
    checked = checked_days_by_employee(df)
    sum_df = sum_df.drop(columns=["วันทำงานทั้งหมด"], errors="ignore").merge(checked, on=["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"], how="left")
    sum_df["วันทำงานทั้งหมด"] = pd.to_numeric(sum_df["วันทำงานทั้งหมด"], errors="coerce").fillna(0).astype(int)
    sum_df = order_summary_columns(sum_df)
    st.dataframe(sum_df, use_container_width=True)
    
    st.subheader("รายละเอียดรายวัน")
    day_df = df.copy()
    day_df["วันที่"] = day_df["วันที่"].dt.strftime("%d/%m/%y")
    day_df["รูปเช็คอิน"] = day_df["รูปเช็คอิน"].apply(img_cell)
    day_df["รูปเช็คเอาท์"] = day_df["รูปเช็คเอาท์"].apply(img_cell)
    day_df["เอกสารลาป่วย"] = day_df["เอกสารลาป่วย"].apply(lambda v: "📎 มีเอกสาร" if str(v).startswith("data:") else "-")
    day_cols = [
        "วันที่", "รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง",
        "สถานะ", "เวลาเข้า", "เวลาออก", "จำนวนชั่วโมง", "จำนวนวัน WOP",
        "OT 1.5(ชม.)", "OT 1(ชม.)", "สายชม", "สายนาที",
        "รูปเช็คอิน", "รูปเช็คเอาท์", "เอกสารลาป่วย"
    ]
    st.write(day_df[day_cols].to_html(escape=False, index=False), unsafe_allow_html=True)

    # ------------------------------
    # จัดการข้อมูลรายวัน: ลบรายการที่ผิด
    # ------------------------------
    st.subheader("🗑️ จัดการข้อมูลรายวัน")
    st.caption("หากเวลา/สถานะ/รูปเช็คอิน-เช็คเอาท์ผิด สามารถลบรายการนั้นแล้วกลับไปบันทึกใหม่ได้")

    # ใช้ index เดิมของ df_rec เพื่อให้ลบถูกแถวใน records.csv
    for row_idx, row in df.iterrows():
        with st.container(border=True):
            r1, r2, r3, r4, r5, r6 = st.columns([1.1, 1.4, 2.0, 1.5, 2.0, 1.0])
            r1.write(row["วันที่"].strftime("%d/%m/%y") if hasattr(row["วันที่"], "strftime") else str(row["วันที่"]))
            r2.write(str(row["รหัสพนักงาน"]))
            r3.write(str(row["ชื่อพนักงาน"]))
            r4.write(str(row["สถานะ"]))
            r5.write(f"เข้า {row['เวลาเข้า']} | ออก {row['เวลาออก']}")
            if r6.button("🗑️ ลบ", key=f"delete_record_{row_idx}"):
                st.session_state["confirm_delete_record"] = int(row_idx)
                st.rerun()

            if st.session_state.get("confirm_delete_record") == int(row_idx):
                st.warning("⚠️ ต้องการลบข้อมูลรายการนี้ใช่หรือไม่? การลบจะนำข้อมูลออกจาก records.csv และไม่สามารถกู้คืนจากในระบบได้ เว้นแต่มีไฟล์สำรอง")
                cy, cn = st.columns(2)
                with cy:
                    if st.button("✅ ยืนยันลบ", key=f"confirm_delete_{row_idx}", type="primary"):
                        df_rec = load_records()
                        if int(row_idx) in df_rec.index:
                            df_rec = df_rec.drop(index=int(row_idx))
                            save_records(df_rec)
                            st.session_state.pop("confirm_delete_record", None)
                            st.success("ลบข้อมูลเรียบร้อยแล้ว ✅")
                            st.rerun()
                        else:
                            st.session_state.pop("confirm_delete_record", None)
                            st.warning("ไม่พบข้อมูลรายการนี้ อาจถูกลบไปแล้ว")
                            st.rerun()
                with cn:
                    if st.button("❌ ยกเลิก", key=f"cancel_delete_{row_idx}"):
                        st.session_state.pop("confirm_delete_record", None)
                        st.rerun()

    day_xl = day_df[day_cols].copy()
    day_xl["รูปเช็คอิน"] = day_xl["รูปเช็คอิน"].apply(lambda x: "มีรูป" if "<img" in str(x) else "-")
    day_xl["รูปเช็คเอาท์"] = day_xl["รูปเช็คเอาท์"].apply(lambda x: "มีรูป" if "<img" in str(x) else "-")
    day_xl["เอกสารลาป่วย"] = day_xl["เอกสารลาป่วย"].apply(lambda x: "มีเอกสาร" if str(x).startswith("📎") or str(x).startswith("data:") else "-")
    out_file = f"รายงาน_{start_d.strftime('%Y%m%d')}_{end_d.strftime('%Y%m%d')}.xlsx"
    with pd.ExcelWriter(out_file, engine="openpyxl") as w:
        sum_df.to_excel(w, sheet_name="สรุปรวม", index=False)
        day_xl.to_excel(w, sheet_name="รายละเอียด", index=False)
    with open(out_file, "rb") as f:
        st.download_button("📥 ดาวน์โหลดรายงาน Excel", f, out_file)

# ==========================================
# 5. หน้าสรุปส่ง HR
# ==========================================
elif menu == "หน้าสรุปส่ง HR":
    st.header("หน้าสรุปส่ง HR")
    default_start, default_end = get_payroll_period()
    st.info("📅 ค่าเริ่มต้นเป็นรอบเงินเดือน 26–25 และ HR สามารถเลือกวันที่เริ่ม/สิ้นสุดเองได้")
    c1, c2 = st.columns(2)
    start_d = c1.date_input("วันที่เริ่มต้น", default_start)
    end_d = c2.date_input("วันที่สิ้นสุด", default_end)
    st.subheader(f"รอบ {start_d.strftime('%d/%m/%y')} - {end_d.strftime('%d/%m/%y')}")
    if start_d > end_d:
        st.error("วันที่เริ่มต้องไม่เกินวันที่สิ้นสุด")
        st.stop()
    df_rec = load_records()
    if df_rec.empty:
        st.info("ยังไม่มีข้อมูล")
        st.stop()
    df_rec["วันที่"] = pd.to_datetime(df_rec["วันที่"], errors="coerce")
    mask = (df_rec["วันที่"] >= pd.to_datetime(start_d)) & (df_rec["วันที่"] <= pd.to_datetime(end_d))
    df = df_rec.loc[mask].copy()
    # เรียงวันที่ตามจริงก่อนสร้างตาราง HR เพื่อให้คอลัมน์วันที่เรียงตามรอบเงินเดือน
    df = df.sort_values(["วันที่", "รหัสพนักงาน"], ascending=[True, True], na_position="last")
    df = add_unpaid_leave_hours(df)
    if df.empty:
        st.info("ไม่มีข้อมูลในช่วงนี้")
        st.stop()
    
    df["วันที่_str"] = df["วันที่"].dt.strftime("%d/%m/%y")
    # ช่องวันที่ในสรุป HR ให้เป็น 1 เมื่อวันนั้นมีเวลาเข้าและเวลาออกจริง
    # ไม่ว่าพนักงานจะเป็นสถานะมาปกติ หรือลากิจไม่รับเงิน(ชม.)
    _tin = df["เวลาเข้า"].astype(str).str.strip()
    _tout = df["เวลาออก"].astype(str).str.strip()
    _has_in_out = (
        _tin.ne("") & _tout.ne("") &
        ~_tin.isin(["-", "nan", "NaT"]) &
        ~_tout.isin(["-", "nan", "NaT"])
    )
    df["ค่า"] = _has_in_out.astype(int)
    pivot = df.pivot_table(
        index=["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"],
        columns="วันที่_str", values="ค่า", aggfunc="first", fill_value=""
    ).reset_index()

    # บังคับลำดับคอลัมน์วันที่ให้เรียงตามวันที่จริง ไม่ใช่เรียงตามข้อความ dd/mm/yy
    date_cols = [c for c in pivot.columns if c not in ["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"]]
    date_cols = sorted(date_cols, key=lambda x: datetime.strptime(str(x), "%d/%m/%y"))
    pivot = pivot[["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"] + date_cols]
    
    emp_summary = df.groupby(
        ["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"], dropna=False
    ).apply(lambda g: pd.Series({
        "มาปกติ(วัน)": g["สถานะ"].isin(["มาปกติ", "ทำงานล่วงเวลา"]).sum(),
        "ลาป่วย(วัน)": (g["สถานะ"] == "ลาป่วย").sum(),
        "WOP(วัน)": pd.to_numeric(g["จำนวนวัน WOP"], errors="coerce").fillna(0).sum(),
        "ลากิจไม่รับเงิน(ชม.)": g["_ลากิจไม่รับเงินชม"].sum(),
        "ขาดงาน(วัน)": (g["สถานะ"] == "ขาดงาน").sum(),
        "OT 1.5(ชม.)": g["OT 1.5(ชม.)"].sum(),
        "OT 1(ชม.)": g["OT 1(ชม.)"].sum(),
        "มาสายรวม(ชม.)": g["สายชม"].sum(),
        "มาสายรวม(นาที)": g["สายนาที"].sum()
    })).reset_index()
    
    final_df = pd.merge(pivot, emp_summary,
                        on=["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"])
    checked = checked_days_by_employee(df)
    final_df = final_df.drop(columns=["วันทำงานทั้งหมด"], errors="ignore").merge(checked, on=["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"], how="left")
    final_df["วันทำงานทั้งหมด"] = pd.to_numeric(final_df["วันทำงานทั้งหมด"], errors="coerce").fillna(0).astype(int)
    final_df = order_hr_columns(final_df)
    st.dataframe(final_df, use_container_width=True)
    
    fn = f"สรุปส่งHR_{start_d.strftime('%Y%m%d')}_{end_d.strftime('%Y%m%d')}.xlsx"
    with pd.ExcelWriter(fn, engine="openpyxl") as w:
        final_df.to_excel(w, index=False, sheet_name="สรุปส่งHR")
    with open(fn, "rb") as f:
        st.download_button("📥 ดาวน์โหลดไฟล์ Excel", f, fn)

# ==========================================
# 6. สำรองข้อมูล
# ==========================================
elif menu == "📥 สำรองข้อมูล":
    st.header("สำรองข้อมูล")
    st.info("💡 CSV เก็บข้อมูลเป็นข้อความ จึงไม่สามารถแสดงรูปภาพได้โดยตรง หากต้องการสำรองรูปเช็คอิน/เช็คเอาท์ ให้ใช้ ZIP หรือ Excel ด้านล่าง")

    # รายชื่อพนักงาน CSV (ภาษาไทยรองรับ Excel)
    if os.path.exists(EMP_FILE):
        emp_download = load_employees()
        emp_csv = emp_download.to_csv(index=False).encode("utf-8-sig")
        st.download_button(
            "📥 ดาวน์โหลด: รายชื่อพนักงาน.csv",
            data=emp_csv,
            file_name="รายชื่อพนักงาน.csv",
            mime="text/csv"
        )

    # ข้อมูลบันทึกเวลา CSV: เก็บข้อมูลข้อความ/รูปแบบ data URI แต่ Excel จะเห็นเป็นข้อความยาว
    if True:
        rec_download = load_records()
        rec_csv = rec_download.to_csv(index=False).encode("utf-8-sig")
        st.download_button(
            "📥 ดาวน์โหลด: ข้อมูลบันทึกเวลา.csv (ข้อมูลข้อความ)",
            data=rec_csv,
            file_name="ข้อมูลบันทึกเวลา.csv",
            mime="text/csv"
        )

        # สำรองพร้อมรูปเป็น ZIP: CSV + ไฟล์รูปจริง
        zip_buf = io.BytesIO()
        image_count = 0
        with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("ข้อมูลบันทึกเวลา.csv", rec_csv)
            for i, row in rec_download.reset_index(drop=True).iterrows():
                for field, prefix in [("รูปเช็คอิน", "เช็คอิน"), ("รูปเช็คเอาท์", "เช็คเอาท์")]:
                    val = row.get(field, "")
                    if pd.isna(val):
                        continue
                    val = str(val)
                    if not val.startswith("data:image"):
                        continue
                    try:
                        header, b64 = val.split(",", 1)
                        ext = "jpg"
                        if "png" in header:
                            ext = "png"
                        elif "webp" in header:
                            ext = "webp"
                        elif "jpeg" in header or "jpg" in header:
                            ext = "jpg"
                        eid = str(row.get("รหัสพนักงาน", ""))
                        d = str(row.get("วันที่", "")).replace("/", "-").replace(" ", "_")
                        filename = f"รูป/{i+1}_{eid}_{d}_{prefix}.{ext}"
                        zf.writestr(filename, base64.b64decode(b64))
                        image_count += 1
                    except Exception:
                        pass
        zip_buf.seek(0)
        st.download_button(
            f"🗂️ ดาวน์โหลดสำรองข้อมูลพร้อมรูป (ZIP) — พบรูป {image_count} รูป",
            data=zip_buf.getvalue(),
            file_name="สำรองข้อมูล_พร้อมรูป.zip",
            mime="application/zip"
        )

        # Excel สำรองพร้อมรูปฝังในเซลล์
        try:
            from openpyxl import Workbook
            from openpyxl.drawing.image import Image as XLImage
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = "ข้อมูลบันทึกเวลา"
            cols = list(rec_download.columns)
            for c, col in enumerate(cols, 1):
                ws.cell(1, c, col)

            image_cols = {"รูปเช็คอิน": 11, "รูปเช็คเอาท์": 12}
            for r_idx, (_, row) in enumerate(rec_download.iterrows(), 2):
                for c_idx, col in enumerate(cols, 1):
                    val = row[col]
                    if col in image_cols and isinstance(val, str) and val.startswith("data:image"):
                        try:
                            _, b64 = val.split(",", 1)
                            img_bytes = base64.b64decode(b64)
                            img = XLImage(io.BytesIO(img_bytes))
                            img.width = 90
                            img.height = 70
                            ws.add_image(img, f"{get_column_letter(c_idx)}{r_idx}")
                            ws.row_dimensions[r_idx].height = 55
                            ws.cell(r_idx, c_idx, "มีรูป")
                        except Exception:
                            ws.cell(r_idx, c_idx, "มีรูป (เปิดไม่ได้)")
                    else:
                        if pd.isna(val):
                            val = ""
                        ws.cell(r_idx, c_idx, str(val))

            for c, col in enumerate(cols, 1):
                ws.column_dimensions[get_column_letter(c)].width = 18
            ws.column_dimensions["K"].width = 16
            ws.column_dimensions["L"].width = 16
            excel_buf = io.BytesIO()
            wb.save(excel_buf)
            excel_buf.seek(0)
            st.download_button(
                "📊 ดาวน์โหลด Excel พร้อมรูปภาพ",
                data=excel_buf.getvalue(),
                file_name="สำรองข้อมูล_พร้อมรูป.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        except Exception as e:
            st.warning(f"ไม่สามารถสร้าง Excel พร้อมรูปได้: {e}")

    st.info("💡 แนะนำ: ถ้าต้องการเก็บรูปไว้ด้วย ให้ดาวน์โหลด ZIP หรือ Excel พร้อมรูป ไม่ควรใช้ CSV เป็นไฟล์สำรองรูปภาพ")
    st.subheader("กู้คืนข้อมูล")
    up_emp = st.file_uploader("อัปโหลด: รายชื่อพนักงาน.csv", type="csv")
    up_rec = st.file_uploader("อัปโหลด: ข้อมูลบันทึกเวลา.csv", type="csv")
    if up_emp and st.button("✅ บันทึกรายชื่อพนักงาน"):
        try:
            try:
                imported_emp = pd.read_csv(up_emp, encoding="utf-8-sig")
            except UnicodeDecodeError:
                imported_emp = pd.read_csv(up_emp, encoding="cp874")
            required_emp = ["รหัสพนักงาน", "ชื่อพนักงาน", "ชื่อเล่น", "ตำแหน่ง"]
            missing_emp = [c for c in required_emp if c not in imported_emp.columns]
            if missing_emp:
                st.error("ไฟล์รายชื่อพนักงานขาดคอลัมน์: " + ", ".join(missing_emp))
            elif save_employees(imported_emp):
                st.success("บันทึกรายชื่อพนักงานสำเร็จ 🎉")
                st.balloons()
                st.rerun()
        except Exception as e:
            st.error(f"บันทึกรายชื่อพนักงานไม่สำเร็จ: {e}")
    if up_rec and st.button("✅ บันทึกข้อมูลบันทึกเวลา"):
        try:
            try:
                imported_rec = pd.read_csv(up_rec, encoding="utf-8-sig")
            except UnicodeDecodeError:
                imported_rec = pd.read_csv(up_rec, encoding="cp874")
            missing_rec = [c for c in REC_COLUMNS if c not in imported_rec.columns]
            if missing_rec:
                st.error("ไฟล์ข้อมูลบันทึกเวลาขาดคอลัมน์: " + ", ".join(missing_rec))
            elif save_records(imported_rec):
                st.success("บันทึกข้อมูลสำเร็จ 🎉")
                st.balloons()
                st.rerun()
        except Exception as e:
            st.error(f"บันทึกข้อมูลบันทึกเวลาไม่สำเร็จ: {e}")

